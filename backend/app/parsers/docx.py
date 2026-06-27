import docx
import re
from typing import List
from uuid import UUID
from app.parsers.base import BaseParser
from app.schemas.schemas import RawRecordCreate
from app.core.logger import logger

class DOCXParser(BaseParser):
    def parse(self, file_path: str, source_file_id: UUID) -> List[RawRecordCreate]:
        logger.info(f"Parsing DOCX file: {file_path}")
        records = []
        try:
            doc = docx.Document(file_path)
            
            # Strategy 1: Read tables
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if len(cells) < 2:
                        continue
                    
                    service_name = cells[0]
                    price_str = cells[-1]
                    
                    # Handle 3+ column table with code (e.g., U1.1 | Service Name | Price)
                    if len(cells) >= 3 and re.match(r'^[A-Z0-9\.\-]{1,15}$', cells[0], re.IGNORECASE):
                        service_name = cells[1]   # Middle column = name
                        price_str = cells[-1]     # Last column = price
                    elif len(cells) == 2:
                        service_name = cells[0]
                        price_str = cells[1]
                    
                    # Truncate to fit DB column limits
                    service_name = service_name[:490]
                    price_str = price_str[:95]
                    
                    # Validate: service_name must look like text, price_str must contain digits
                    if len(service_name) > 3 and any(c.isalpha() for c in service_name) and re.search(r'\d', price_str):
                        # Skip header rows
                        if service_name.lower() in ('наименование', 'услуга', 'название', 'service'):
                            continue
                        records.append(RawRecordCreate(
                            source_file_id=source_file_id,
                            raw_service_name=service_name,
                            raw_price=price_str
                        ))
                            
            # Strategy 2: Read paragraphs
            if len(records) < 5:
                for para in doc.paragraphs:
                    line = para.text.strip()
                    if not line:
                        continue
                    match = re.search(r'(.*?)\s+([\d\s\.,]+)\s*(?:KZT|тнг|тенге|₸|руб)?$', line, re.IGNORECASE)
                    if match:
                        service_name = match.group(1).strip()
                        price_str = match.group(2).strip()
                        if len(service_name) > 3 and any(c.isalpha() for c in service_name):
                            if not any(r.raw_service_name == service_name for r in records):
                                records.append(RawRecordCreate(
                                    source_file_id=source_file_id,
                                    raw_service_name=service_name,
                                    raw_price=price_str
                                ))

            logger.info(f"Successfully parsed {len(records)} records from DOCX")
        except Exception as e:
            logger.error(f"Failed to parse DOCX file {file_path}: {str(e)}")
            raise e
        return records
